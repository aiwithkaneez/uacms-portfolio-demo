from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def build_xlsx(headers: list[str], rows: list[list[object]], sheet_title: str = "Complaints") -> bytes:
    """Build an .xlsx file in memory — no temp files, the caller streams the
    returned bytes straight back as the HTTP response body.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]  # Excel's own sheet-name length limit

    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append(row)

    for i, header in enumerate(headers, start=1):
        col_letter = get_column_letter(i)
        longest = max([len(header)] + [len(str(row[i - 1])) for row in rows]) if rows else len(header)
        ws.column_dimensions[col_letter].width = min(60, max(10, longest + 2))

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
